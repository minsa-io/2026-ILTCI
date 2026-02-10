"""Rich text formatting for PowerPoint paragraphs."""

import re
from pptx.oxml.xmlchemy import OxmlElement
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from pptx.text.text import _Paragraph

# EMU conversions (914400 EMU = 1 inch)
# Bullet hanging indent: text at 0.5", bullet character at 0"
BULLET_MARGIN_EMU = 457200        # 0.5 inches - where text starts
BULLET_INDENT_EMU = -457200       # -0.5 inches - bullet hangs to the left
SUB_BULLET_MARGIN_EMU = 914400    # 1.0 inches - where sub-bullet text starts
SUB_BULLET_INDENT_EMU = -457200   # -0.5 inches - sub-bullet hangs to the left


def _set_paragraph_margins(pPr, marL: int, indent: int) -> None:
    """Set paragraph margins via XML attributes.
    
    Args:
        pPr: Paragraph properties element
        marL: Left margin in EMU
        indent: First line indent in EMU (negative for hanging)
    """
    pPr.set('marL', str(marL))
    pPr.set('indent', str(indent))


def add_bullet(paragraph: '_Paragraph', level: int = 0) -> None:
    """Add bullet formatting to a paragraph with proper hanging indent.
    
    Args:
        paragraph: PowerPoint paragraph object
        level: Indentation level (0 = first level bullet, 1 = sub-bullet)
    """
    pPr = paragraph._element.get_or_add_pPr()
    
    # Set hanging indent based on level
    # Level 0: marL=0.5", indent=-0.5" (text at 0.5", bullet at 0")
    # Level 1: marL=1.0", indent=-0.5" (text at 1.0", bullet at 0.5")
    if level == 0:
        _set_paragraph_margins(pPr, BULLET_MARGIN_EMU, BULLET_INDENT_EMU)
    else:
        _set_paragraph_margins(pPr, SUB_BULLET_MARGIN_EMU, SUB_BULLET_INDENT_EMU)
    
    # Create bullet element
    buChar = OxmlElement('a:buChar')
    buChar.set('char', '•')  # Bullet character
    pPr.insert(0, buChar)


def remove_bullet(paragraph: '_Paragraph') -> None:
    """Remove bullet formatting from a paragraph and reset margins to zero.
    
    This ensures plain text paragraphs don't inherit bullet margins from lstStyle.
    
    Args:
        paragraph: PowerPoint paragraph object
    """
    pPr = paragraph._element.get_or_add_pPr()
    
    # Set margins to zero for plain text
    _set_paragraph_margins(pPr, 0, 0)
    
    pPr.insert(0, OxmlElement('a:buNone'))


def add_numbering(paragraph: '_Paragraph', start_at: int = 1, numbering_type: str = 'arabicPeriod') -> None:
    """Add automatic numbering to a paragraph with proper hanging indent.
    
    Args:
        paragraph: PowerPoint paragraph object
        start_at: Starting number
        numbering_type: Numbering style (e.g., 'arabicPeriod' for 1. 2. 3.)
    """
    pPr = paragraph._element.get_or_add_pPr()
    
    # Set hanging indent for numbered items (same as level-0 bullets)
    _set_paragraph_margins(pPr, BULLET_MARGIN_EMU, BULLET_INDENT_EMU)
    
    # Create buAutoNum element for numbering
    buAutoNum = OxmlElement('a:buAutoNum')
    buAutoNum.set('type', numbering_type)
    if start_at > 1:
        buAutoNum.set('startAt', str(start_at))
    pPr.insert(0, buAutoNum)


def add_formatted_text(paragraph: '_Paragraph', text: str) -> None:
    """Add text to a paragraph with markdown formatting support.
    
    Supports **bold**, *italic*, ***bold+italic***, and [text](url) markdown syntax.
    
    Args:
        paragraph: PowerPoint paragraph object
        text: Text with markdown formatting
    """
    # Clear any existing text
    paragraph.text = ""
    
    # Pattern to match markdown formatting including links
    # Matches [text](url) links, ***text*** (bold+italic), **text** (bold), or *text* (italic)
    pattern = r'(\[.*?\]\(.*?\)|\*\*\*.*?\*\*\*|\*\*.*?\*\*|\*.*?\*)'
    
    # Split text by markdown patterns
    parts = re.split(pattern, text)
    
    for part in parts:
        if not part:
            continue
        
        # Check what type of formatting this part has
        if part.startswith('[') and part.endswith(')') and '](' in part:
            # Markdown link [text](url)
            # Extract the link text and URL
            link_match = re.match(r'\[(.*?)\]\((.*?)\)', part)
            if link_match:
                link_text = link_match.group(1)
                link_url = link_match.group(2)
                run = paragraph.add_run()
                run.text = link_text
                # Add hyperlink
                run.hyperlink.address = link_url
                # Add underline for visibility (PowerPoint will handle color)
                run.font.underline = True
        elif part.startswith('***') and part.endswith('***'):
            # Bold and italic
            run = paragraph.add_run()
            run.text = part[3:-3]
            run.font.bold = True
            run.font.italic = True
        elif part.startswith('**') and part.endswith('**'):
            # Bold
            run = paragraph.add_run()
            run.text = part[2:-2]
            run.font.bold = True
        elif part.startswith('*') and part.endswith('*'):
            # Italic
            run = paragraph.add_run()
            run.text = part[1:-1]
            run.font.italic = True
        else:
            # Regular text
            run = paragraph.add_run()
            run.text = part
